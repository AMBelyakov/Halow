# -*- coding: utf-8 -*-
"""Общая вёрстка разделов ВКР (ГОСТ-подобное оформление) на python-docx.

Формулы — нативные формулы Word (OMML). Ссылки на источники пишутся в тексте
как «[@ключ]» или «[@a; @b]» и нумеруются при сохранении по порядку первого
упоминания; список источников строится из словаря ключ -> описание.
"""
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

NBSP = " "
TEXT_W = 165  # мм: A4 210 - 30 - 15
LEFT, CENTER, RIGHT = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT


def n(x, d=0):
    """Число по-русски: неразрывные пробелы между тысячами, запятая."""
    return f"{x:,.{d}f}".replace(",", NBSP).replace(".", ",")


UNITS = (" %", " руб.", " чел.-ч", " кал. дн", " раб. дн", " кВт", " Вт", " лк", " дБА", " дБ",
         " м²", " м³", " м/с", " °C", " мм", " см", " м;", " м.", " м,", " м ", " ч;", " ч.",
         " ч,", " ч ", " г.", " года", " лет", " месяцев", " долл.", " мкВт/см²", " лм", " МГц", " ГГц")


PREFIXES = ("№ ", "ГОСТ ", "ГОСТ Р ", "СП ", "СанПиН ", "МР ", "Р 2.", "ОУ-", "ОП-")


def nb(s):
    """Неразрывный пробел между числом и единицей измерения, после «№», «ГОСТ» и т. п."""
    for u in UNITS:
        s = s.replace(u, NBSP + u[1:])
    for pfx in PREFIXES:
        s = s.replace(pfx, pfx.replace(" ", NBSP))
    return s


# ============================================================== OMML
MNS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def mr(text, plain=False):
    rpr = '<m:rPr><m:sty m:val="p"/></m:rPr>' if plain else ""
    return (f'<m:r>{rpr}<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/></w:rPr>'
            f'<m:t xml:space="preserve">{escape(text)}</m:t></m:r>')


def V(sym, sub=None):
    base = mr(sym)
    if sub is None:
        return base
    return f"<m:sSub><m:e>{base}</m:e><m:sub>{mr(sub, True)}</m:sub></m:sSub>"


def P(text):
    return mr(text, True)


def sup(base, s):
    return f"<m:sSup><m:e>{base}</m:e><m:sup>{s}</m:sup></m:sSup>"


def frac(num, den):
    return f"<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>"


def msum(e, idx=None, lo=None, hi=None):
    if lo is not None:
        return (f'<m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>'
                f"<m:sub>{lo}</m:sub><m:sup>{hi}</m:sup><m:e>{e}</m:e></m:nary>")
    if idx is None:
        pr, sub = '<m:subHide m:val="1"/><m:supHide m:val="1"/>', "<m:sub/>"
    else:
        pr, sub = '<m:supHide m:val="1"/>', f"<m:sub>{mr(idx)}</m:sub>"
    return (f'<m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/>{pr}</m:naryPr>'
            f"{sub}<m:sup/><m:e>{e}</m:e></m:nary>")


def bar(e):
    return f'<m:bar><m:barPr><m:pos m:val="top"/></m:barPr><m:e>{e}</m:e></m:bar>'


def paren(e):
    return f"<m:d><m:e>{e}</m:e></m:d>"


def omath(*parts):
    return parse_xml(f'<m:oMath xmlns:m="{MNS}" xmlns:w="{WNS}">' + "".join(parts) + "</m:oMath>")


def _set_font(style, size, bold=False):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), "Times New Roman")


class VkrDoc:
    """Раздел ВКР: главa `chapter`, источники с префиксом `ref_prefix`."""

    def __init__(self, chapter, ref_prefix):
        self.ch, self.pref = chapter, ref_prefix
        self.fcount = self.tcount = self.picount = 0
        d = self.doc = Document()
        sec = d.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin, sec.right_margin = Mm(30), Mm(15)
        sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)
        self.sec = sec

        normal = d.styles["Normal"]
        _set_font(normal, 14)
        pf = normal.paragraph_format
        pf.first_line_indent = Cm(1.25)
        pf.line_spacing = 1.5
        pf.space_before = pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.widow_control = True

        for lvl, size in ((1, 16), (2, 14), (3, 14)):
            st = d.styles[f"Heading {lvl}"]
            _set_font(st, size, bold=True)
            hp = st.paragraph_format
            hp.alignment = CENTER
            hp.first_line_indent = Cm(0)
            hp.line_spacing = 1.5
            hp.space_before, hp.space_after = Pt(12), Pt(6)
            hp.keep_with_next = True

        tt = d.styles.add_style("TableText", WD_STYLE_TYPE.PARAGRAPH)
        tt.base_style = normal
        _set_font(tt, 12)
        tt.paragraph_format.first_line_indent = Cm(0)
        tt.paragraph_format.line_spacing = 1.0
        tt.paragraph_format.alignment = LEFT

        cap = d.styles.add_style("TableCaption", WD_STYLE_TYPE.PARAGRAPH)
        cap.base_style = normal
        cap.paragraph_format.first_line_indent = Cm(0)
        cap.paragraph_format.alignment = LEFT
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(6)

        fs = d.styles.add_style("Formula", WD_STYLE_TYPE.PARAGRAPH)
        fs.base_style = normal
        fs.paragraph_format.first_line_indent = Cm(0)
        fs.paragraph_format.alignment = LEFT
        fs.paragraph_format.space_before = fs.paragraph_format.space_after = Pt(6)

        self._footer_page_numbers()

    # ---------------------------------------------------------- текст
    def H(self, text, level):
        return self.doc.add_heading(text, level)

    def para(self, *parts, indent=True, align=None, keep=False):
        """parts: str | ('m', oMath) | ('b', str) | ('h', str — выделить жёлтым)."""
        p = self.doc.add_paragraph()
        if not indent:
            p.paragraph_format.first_line_indent = Cm(0)
        if align is not None:
            p.paragraph_format.alignment = align
        if keep:
            p.paragraph_format.keep_with_next = True
        for part in parts:
            if isinstance(part, str):
                p.add_run(nb(part))
            elif part[0] == "m":
                p._p.append(part[1])
            elif part[0] == "b":
                p.add_run(nb(part[1])).bold = True
            elif part[0] == "h":
                p.add_run(nb(part[1])).font.highlight_color = WD_COLOR_INDEX.YELLOW
        return p

    @staticmethod
    def m(*parts):
        return ("m", omath(*parts))

    def dash(self, *parts):
        return self.para("– ", *parts)

    def where(self, *parts):
        return self.para("где ", *parts, indent=False)

    def cont(self, *parts):
        return self.para(*parts)

    # ---------------------------------------------------------- формулы
    @staticmethod
    def _no_borders(t):
        b = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), "nil")
            b.append(e)
        t._tbl.tblPr.append(b)

    def formula(self, *parts, numbered=True):
        """Нумерованная формула — в невидимой таблице (формула рисуется в выключном
        режиме, номер справа); ненумерованная — отдельный абзац по центру."""
        if not numbered:
            p = self.doc.add_paragraph(style="Formula")
            p.paragraph_format.alignment = CENTER
            p._p.append(omath(*parts))
            return None
        self.fcount += 1
        num = f"({self.ch}.{self.fcount})"
        widths = (18, 129, 18)
        t = self.doc.add_table(rows=1, cols=3)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        self._no_borders(t)
        for ci, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths[ci] / 25.4 * 1440)))
        for ci, cell in enumerate(t.rows[0].cells):
            cell.width = Mm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.style = self.doc.styles["Formula"]
            if ci == 1:
                p.paragraph_format.alignment = CENTER
                p._p.append(omath(*parts))
            elif ci == 2:
                p.paragraph_format.alignment = RIGHT
                p.add_run(num)
        return num

    # ---------------------------------------------------------- таблицы
    def table(self, caption, widths_mm, rows, header_rows=1, size=12, hsize=None, align=None,
              merges=(), bold_rows=()):
        """rows: список строк; ячейка — str или ('h', str) для выделения жёлтым."""
        self.tcount += 1
        num = f"{self.ch}.{self.tcount}"
        self.para(f"Таблица {num} – {caption}", indent=False).style = self.doc.styles["TableCaption"]
        t = self.doc.add_table(rows=len(rows), cols=len(widths_mm))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        mar = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), "57")
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        t._tbl.tblPr.append(mar)
        al = {"l": LEFT, "c": CENTER, "r": RIGHT}
        for ri, row in enumerate(rows):
            tr = t.rows[ri]
            trPr = tr._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if ri < header_rows:
                th = OxmlElement("w:tblHeader")
                th.set(qn("w:val"), "true")
                trPr.append(th)
            for ci, val in enumerate(row):
                hl = isinstance(val, tuple) and val[0] == "h"
                text = val[1] if hl else str(val)
                cell = tr.cells[ci]
                cell.width = Mm(widths_mm[ci])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                for li, line in enumerate(text.split("\n")):
                    if li:
                        p = cell.add_paragraph()
                    p.style = self.doc.styles["TableText"]
                    run = p.add_run(line)
                    run.font.size = Pt(hsize if (ri < header_rows and hsize) else size)
                    if hl:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    if ri < header_rows or ri in bold_rows:
                        run.bold = True
                    if ri < header_rows:
                        p.alignment = CENTER
                    elif align:
                        p.alignment = al[align[ci]]
        for ci, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths_mm[ci] / 25.4 * 1440)))
        if len(rows) <= 15:                       # небольшая таблица не разрывается
            for tr in t.rows[:-1]:
                for cell in tr.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
        for (r0, c0, r1, c1) in merges:
            merged = t.cell(r0, c0).merge(t.cell(r1, c1))
            for p in merged.paragraphs[1:]:
                p._p.getparent().remove(p._p)
        self.doc.add_paragraph().paragraph_format.line_spacing = 1.0
        return num

    # ---------------------------------------------------------- рисунки
    def figure(self, path, caption, width_mm=TEXT_W):
        self.picount += 1
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.alignment = CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Mm(width_mm))
        self.para(f"Рисунок {self.ch}.{self.picount} – {caption}", indent=False, align=CENTER)
        return f"{self.ch}.{self.picount}"

    def next_fig(self):
        return f"{self.ch}.{self.picount + 1}"

    def next_tab(self):
        return f"{self.ch}.{self.tcount + 1}"

    # ---------------------------------------------------------- колонтитул
    def _footer_page_numbers(self):
        footer = self.sec.footer.paragraphs[0]
        footer.alignment = CENTER
        run = footer.add_run()
        for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = txt
            run._r.append(el)

    # ---------------------------------------------------------- источники
    def finalize(self, sources, title, out):
        """Нумерует ссылки [@key] по первому упоминанию, добавляет список, сохраняет."""
        order = {}
        tok = re.compile(r"\[@([^\]]+)\]")

        def repl(mo):
            keys = [k.strip().lstrip("@") for k in mo.group(1).split(";")]
            nums = []
            for k in keys:
                if k not in sources:
                    raise KeyError(f"нет источника «{k}»")
                order.setdefault(k, len(order) + 1)
                nums.append(order[k])
            return "[" + "; ".join(f"{self.pref}{x}" for x in nums) + "]"

        for t in self.doc.element.body.iter(qn("w:t")):
            if t.text and "[@" in t.text:
                t.text = tok.sub(repl, t.text)
        unused = set(sources) - set(order)
        self.H(title, 2)
        self.para(f"(номера [{self.pref}1]–[{self.pref}{len(order)}] привести к сквозной нумерации "
                  "общего списка использованных источников)", indent=False, align=CENTER)
        for k, num in sorted(order.items(), key=lambda kv: kv[1]):
            self.para(f"[{self.pref}{num}] {sources[k]}", indent=False, align=LEFT)
        self.doc.save(out)
        return order, unused
